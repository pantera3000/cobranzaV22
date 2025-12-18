from django.shortcuts import render, get_object_or_404, redirect
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse
from django.contrib import messages
from .models import Cobrador
from .forms import CobradorForm
import csv
from openpyxl import Workbook
from devoluciones.models import Devolucion  # ✅ Asegúrate de importar
from datetime import timedelta
from django.db.models import Sum, F
from django.utils import timezone
from cobros.models import Cobro
from documentos.models import Documento
from django.core.paginator import Paginator
from django.utils import timezone
import calendar
from datetime import datetime


from django.contrib.auth.decorators import user_passes_test

from openpyxl.styles import Font











def es_admin_o_elegido(user):
    return user.username in ['adminaqp', 'juanc', 'maria']  # 👈 Cambia por los que desees





def localtime_peru():
    return timezone.localtime(timezone.now())



def get_filtro_label(filtro_rapido, fecha_desde, fecha_hasta):
    """Devuelve una etiqueta descriptiva del filtro aplicado"""
    hoy = localtime_peru().date()

    if filtro_rapido == 'hoy':
        return f"Hoy – {hoy.strftime('%d de %B de %Y')}"
    elif filtro_rapido == 'ayer':
        ayer = hoy - timedelta(days=1)
        return f"Ayer – {ayer.strftime('%d de %B de %Y')}"
    elif filtro_rapido == 'mes':
        inicio = hoy.replace(day=1)
        return f"Este mes – {inicio.strftime('%d de %B')} al {hoy.strftime('%d de %B de %Y')}"
    elif filtro_rapido == 'mes_pasado':
        ultimo_dia = hoy.replace(day=1) - timedelta(days=1)
        primer_dia = ultimo_dia.replace(day=1)
        return f"Mes pasado – {primer_dia.strftime('%d de %B')} al {ultimo_dia.strftime('%d de %B de %Y')}"
    elif filtro_rapido == '3meses':
        inicio = hoy - timedelta(days=90)
        return f"Últimos 3 meses – {inicio.strftime('%d de %B')} al {hoy.strftime('%d de %B de %Y')}"
    elif filtro_rapido == 'año':
        inicio = hoy.replace(month=1, day=1)
        return f"Este año – {inicio.strftime('%Y')} (desde {inicio.strftime('%d de %B')})"
    elif filtro_rapido == 'año_pasado':
        año = hoy.year - 1
        return f"Año pasado – {año}"
    else:
        try:
            from django.utils.dateparse import parse_date
            fecha_desde_dt = parse_date(fecha_desde)
            fecha_hasta_dt = parse_date(fecha_hasta)
            if fecha_desde_dt and fecha_hasta_dt:
                return f"Personalizado – {fecha_desde_dt.strftime('%d/%m')} al {fecha_hasta_dt.strftime('%d/%m/%Y')}"
            elif fecha_desde_dt:
                return f"Desde – {fecha_desde_dt.strftime('%d/%m/%Y')}"
            elif fecha_hasta_dt:
                return f"Hasta – {fecha_hasta_dt.strftime('%d/%m/%Y')}"
        except:
            pass
        return "Filtro personalizado (sin fechas definidas)"




def cobrador_detail(request, pk):
    cobrador = get_object_or_404(Cobrador, pk=pk)
    hoy = localtime_peru().date()

    # === Filtro de fecha GLOBAL ===
    filtro_rapido = request.GET.get('filtro')  # 'hoy', 'ayer', etc.
    fecha_desde = request.GET.get('fecha_desde')
    fecha_hasta = request.GET.get('fecha_hasta')

    # Calcular rango según filtro rápido
    if filtro_rapido == 'hoy':
        fecha_desde = fecha_hasta = hoy.isoformat()
    elif filtro_rapido == 'ayer':
        ayer = hoy - timedelta(days=1)
        fecha_desde = fecha_hasta = ayer.isoformat()
    elif filtro_rapido == 'mes':
        fecha_desde = (hoy.replace(day=1)).isoformat()
        fecha_hasta = hoy.isoformat()
    elif filtro_rapido == 'mes_pasado':
        primer_dia = (hoy.replace(day=1) - timedelta(days=1)).replace(day=1)
        ultimo_dia = hoy.replace(day=1) - timedelta(days=1)
        fecha_desde = primer_dia.isoformat()
        fecha_hasta = ultimo_dia.isoformat()
    elif filtro_rapido == '3meses':
        fecha_desde = (hoy - timedelta(days=90)).isoformat()
        fecha_hasta = hoy.isoformat()
    elif filtro_rapido == 'año':
        fecha_desde = hoy.replace(month=1, day=1).isoformat()
        fecha_hasta = hoy.isoformat()
    elif filtro_rapido == 'año_pasado':
        año_pasado = hoy.year - 1
        fecha_desde = f"{año_pasado}-01-01"
        fecha_hasta = f"{año_pasado}-12-31"

    # ✅ Valores por defecto: Este año
    if not fecha_desde and not fecha_hasta and not filtro_rapido:
        inicio_año = hoy.replace(month=1, day=1)
        fecha_desde = inicio_año.isoformat()
        fecha_hasta = hoy.isoformat()
        filtro_rapido = 'año'

    else:
        if not fecha_desde:
            fecha_desde = (hoy - timedelta(days=30)).isoformat()
        if not fecha_hasta:
            fecha_hasta = hoy.isoformat()

    # Convertir a datetime
    from django.utils.dateparse import parse_date
    fecha_desde_dt = parse_date(fecha_desde) or (hoy - timedelta(days=30))
    fecha_hasta_dt = parse_date(fecha_hasta) or hoy

    # === Datos filtrados por fecha ===
    cobros = Cobro.objects.filter(
        cobrador=cobrador,
        fecha__date__gte=fecha_desde_dt,
        fecha__date__lte=fecha_hasta_dt
    ).select_related('documento', 'documento__cliente').order_by('-fecha')
    cobros_paginator = Paginator(cobros, 20)
    cobros_page = request.GET.get('cobros_page')
    cobros_page_obj = cobros_paginator.get_page(cobros_page)

    devoluciones = Devolucion.objects.filter(
        cobrador=cobrador,
        fecha__date__gte=fecha_desde_dt,
        fecha__date__lte=fecha_hasta_dt
    ).select_related('documento', 'documento__cliente').order_by('-fecha')
    devoluciones_paginator = Paginator(devoluciones, 20)
    devoluciones_page = request.GET.get('devoluciones_page')
    devoluciones_page_obj = devoluciones_paginator.get_page(devoluciones_page)

    # === Documentos Pendientes (con filtro de fecha) ===
    documentos_pendientes = Documento.objects.filter(
        cobrador=cobrador,
        monto_total__gt=F('monto_pagado') + F('monto_devolucion'),
        fecha_emision__date__gte=fecha_desde_dt,  # ✅ Filtro por fecha de emisión
        fecha_emision__date__lte=fecha_hasta_dt
    ).select_related('cliente').order_by('fecha_vencimiento')
    pendientes_paginator = Paginator(documentos_pendientes, 20)
    pendientes_page = request.GET.get('pendientes_page')
    pendientes_page_obj = pendientes_paginator.get_page(pendientes_page)

    documentos_vencidos = documentos_pendientes.filter(
        fecha_vencimiento__lt=timezone.now()
    )
    vencidos_paginator = Paginator(documentos_vencidos, 20)
    vencidos_page = request.GET.get('vencidos_page')
    vencidos_page_obj = vencidos_paginator.get_page(vencidos_page)

    documentos_asignados = Documento.objects.filter(
        cobrador=cobrador,
        fecha_emision__date__gte=fecha_desde_dt,
        fecha_emision__date__lte=fecha_hasta_dt
    ).select_related('cliente').order_by('-fecha_emision')
    documentos_paginator = Paginator(documentos_asignados, 20)
    documentos_page = request.GET.get('documentos_page')
    documentos_page_obj = documentos_paginator.get_page(documentos_page)

    # Cálculos
    total_cobrado = cobros.aggregate(total=Sum('monto'))['total'] or 0
    total_devuelto = devoluciones.aggregate(total=Sum('monto'))['total'] or 0
    total_pendiente = sum(doc.get_saldo_pendiente() for doc in documentos_pendientes)
    total_vencido = sum(doc.get_saldo_pendiente() for doc in documentos_vencidos)

    # ✅ Generar etiqueta del filtro
    filtro_label = get_filtro_label(filtro_rapido, fecha_desde, fecha_hasta)

    # ✅ Calcular cuántos documentos tiene cada referencia
    referencia_count = {}
    for cobro in cobros:
        ref = cobro.referencia
        if ref:
            if ref not in referencia_count:
                referencia_count[ref] = 0
            referencia_count[ref] += 1

    # ✅ Aporte a la Meta Mensual
    try:
        from cobros.models import MetaMensual
        mes_actual = hoy.replace(day=1)
        meta = MetaMensual.objects.get(mes=mes_actual)

        # Fechas del mes
        _, last_day = calendar.monthrange(hoy.year, hoy.month)
        fecha_inicio = timezone.make_aware(datetime.combine(hoy.replace(day=1), datetime.min.time()))
        fecha_fin = timezone.make_aware(datetime.combine(hoy.replace(day=last_day), datetime.max.time()))

        # Cobros del cobrador en el mes
        cobros_mes = Cobro.objects.filter(
            cobrador=cobrador,
            fecha__gte=fecha_inicio,
            fecha__lte=fecha_fin
        ).aggregate(total=Sum('monto'))['total'] or 0

        aporte_cobrador = cobros_mes
        porcentaje_aporte = (aporte_cobrador / meta.monto_objetivo * 100) if meta.monto_objetivo > 0 else 0
        meta_alcanzada = aporte_cobrador >= meta.monto_objetivo

    except Exception as e:
        meta = None
        aporte_cobrador = 0
        porcentaje_aporte = 0
        meta_alcanzada = False

    return render(request, 'cobradores/cobrador_detail.html', {
        'cobrador': cobrador,
        'cobros_page_obj': cobros_page_obj,
        'devoluciones_page_obj': devoluciones_page_obj,
        'pendientes_page_obj': pendientes_page_obj,
        'vencidos_page_obj': vencidos_page_obj,
        'documentos_page_obj': documentos_page_obj,
        'total_cobrado': total_cobrado,
        'total_devuelto': total_devuelto,
        'total_pendiente': total_pendiente,
        'total_vencido': total_vencido,
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'filtro_rapido': filtro_rapido,
        'filtro_label': filtro_label,
        'referencia_count': referencia_count,

        # ✅ Datos de la meta mensual
        'meta': meta,
        'aporte_cobrador': aporte_cobrador,
        'porcentaje_aporte': porcentaje_aporte,
        'meta_alcanzada': meta_alcanzada,
    })











def exportar_cobros_excel(request, pk):
    """
    Exporta SOLO los cobros del cobrador, aplicando los mismos filtros de fecha
    que se usan en cobrador_detail.
    """
    cobrador = get_object_or_404(Cobrador, pk=pk)
    hoy = localtime_peru().date()

    # === Copiar lógica de filtros de cobrador_detail ===
    filtro_rapido = request.GET.get('filtro')
    fecha_desde = request.GET.get('fecha_desde')
    fecha_hasta = request.GET.get('fecha_hasta')

    if filtro_rapido == 'hoy':
        fecha_desde = fecha_hasta = hoy.isoformat()
    elif filtro_rapido == 'ayer':
        ayer = hoy - timedelta(days=1)
        fecha_desde = fecha_hasta = ayer.isoformat()
    elif filtro_rapido == 'mes':
        fecha_desde = (hoy.replace(day=1)).isoformat()
        fecha_hasta = hoy.isoformat()
    elif filtro_rapido == 'mes_pasado':
        primer_dia = (hoy.replace(day=1) - timedelta(days=1)).replace(day=1)
        ultimo_dia = hoy.replace(day=1) - timedelta(days=1)
        fecha_desde = primer_dia.isoformat()
        fecha_hasta = ultimo_dia.isoformat()
    elif filtro_rapido == '3meses':
        fecha_desde = (hoy - timedelta(days=90)).isoformat()
        fecha_hasta = hoy.isoformat()
    elif filtro_rapido == 'año':
        fecha_desde = hoy.replace(month=1, day=1).isoformat()
        fecha_hasta = hoy.isoformat()
    elif filtro_rapido == 'año_pasado':
        año_pasado = hoy.year - 1
        fecha_desde = f"{año_pasado}-01-01"
        fecha_hasta = f"{año_pasado}-12-31"

    if not fecha_desde and not fecha_hasta and not filtro_rapido:
        inicio_año = hoy.replace(month=1, day=1)
        fecha_desde = inicio_año.isoformat()
        fecha_hasta = hoy.isoformat()

    from django.utils.dateparse import parse_date
    fecha_desde_dt = parse_date(fecha_desde) or (hoy - timedelta(days=30))
    fecha_hasta_dt = parse_date(fecha_hasta) or hoy

    # === Obtener cobros filtrados ===
    cobros = Cobro.objects.filter(
        cobrador=cobrador,
        fecha__date__gte=fecha_desde_dt,
        fecha__date__lte=fecha_hasta_dt
    ).select_related('documento', 'documento__cliente').order_by('-fecha')

    # === Crear archivo Excel ===
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cobros"

    # Estilo negrita
    bold_font = Font(bold=True)

    # Encabezados
    headers = ['Documento', 'Cliente', 'Monto', 'Referencia', 'Notas', 'Fecha Pago', 'Registrado por']
    for col_num, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col_num)
        cell.value = header
        cell.font = bold_font

    # Datos
    for cobro in cobros:
        row = [
            f"{cobro.documento.get_tipo_display()} {cobro.documento.get_numero_completo()}",
            cobro.documento.cliente.nombre,
            float(cobro.monto),
            cobro.referencia or '',
            cobro.notas or '',
            cobro.fecha.strftime('%d/%m/%Y %H:%M'),
            cobro.usuario_registro.get_full_name() if cobro.usuario_registro else 'Sistema'
        ]
        sheet.append(row)

    # Ajustar ancho de columnas
    sheet.column_dimensions['A'].width = 20
    sheet.column_dimensions['B'].width = 25
    sheet.column_dimensions['C'].width = 12
    sheet.column_dimensions['D'].width = 18
    sheet.column_dimensions['E'].width = 25
    sheet.column_dimensions['F'].width = 18
    sheet.column_dimensions['G'].width = 20

    # Respuesta HTTP
    filename = f"pagos_{cobrador.nombre}_{timezone.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response





from django.http import HttpResponse
from django.template.loader import render_to_string
from weasyprint import HTML
import tempfile

def exportar_cobros_pdf(request, pk):
    cobrador = get_object_or_404(Cobrador, pk=pk)
    hoy = localtime_peru().date()

    # === Copiar lógica de filtros ===
    filtro_rapido = request.GET.get('filtro')
    fecha_desde = request.GET.get('fecha_desde')
    fecha_hasta = request.GET.get('fecha_hasta')

    if filtro_rapido == 'hoy':
        fecha_desde = fecha_hasta = hoy.isoformat()
    elif filtro_rapido == 'ayer':
        ayer = hoy - timedelta(days=1)
        fecha_desde = fecha_hasta = ayer.isoformat()
    elif filtro_rapido == 'mes':
        fecha_desde = (hoy.replace(day=1)).isoformat()
        fecha_hasta = hoy.isoformat()
    elif filtro_rapido == 'mes_pasado':
        primer_dia = (hoy.replace(day=1) - timedelta(days=1)).replace(day=1)
        ultimo_dia = hoy.replace(day=1) - timedelta(days=1)
        fecha_desde = primer_dia.isoformat()
        fecha_hasta = ultimo_dia.isoformat()
    elif filtro_rapido == '3meses':
        fecha_desde = (hoy - timedelta(days=90)).isoformat()
        fecha_hasta = hoy.isoformat()
    elif filtro_rapido == 'año':
        fecha_desde = hoy.replace(month=1, day=1).isoformat()
        fecha_hasta = hoy.isoformat()
    elif filtro_rapido == 'año_pasado':
        año_pasado = hoy.year - 1
        fecha_desde = f"{año_pasado}-01-01"
        fecha_hasta = f"{año_pasado}-12-31"

    if not fecha_desde and not fecha_hasta and not filtro_rapido:
        inicio_año = hoy.replace(month=1, day=1)
        fecha_desde = inicio_año.isoformat()
        fecha_hasta = hoy.isoformat()

    from django.utils.dateparse import parse_date
    fecha_desde_dt = parse_date(fecha_desde) or (hoy - timedelta(days=30))
    fecha_hasta_dt = parse_date(fecha_hasta) or hoy

    # === Obtener cobros filtrados ===
    cobros = Cobro.objects.filter(
        cobrador=cobrador,
        fecha__date__gte=fecha_desde_dt,
        fecha__date__lte=fecha_hasta_dt
    ).select_related('documento', 'documento__cliente').order_by('-fecha')

    total_cobrado = cobros.aggregate(total=Sum('monto'))['total'] or 0

    # === Generar HTML para PDF ===
    html_string = render_to_string('cobros/cobro_list_pdf.html', {
        'cobrador': cobrador,
        'cobros': cobros,
        'total_cobrado': total_cobrado,
        'fecha_desde': fecha_desde_dt,
        'fecha_hasta': fecha_hasta_dt,
        'filtro_label': get_filtro_label(filtro_rapido, fecha_desde, fecha_hasta),
        'request': request  # Necesario para estáticos
    })

    # === Generar PDF ===
    html = HTML(string=html_string, base_url=request.build_absolute_uri('/'))
    pdf = html.write_pdf()

    # === Preparar respuesta ===
    filename = f"pagos_{cobrador.nombre}_{timezone.now().strftime('%Y%m%d_%H%M')}.pdf"
    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response












































def cobrador_list(request):
    query = request.GET.get('q')
    if query:
        cobradores = Cobrador.objects.filter(
            Q(nombre__icontains=query) | Q(dni__icontains=query)
        )
    else:
        cobradores = Cobrador.objects.all()

    paginator = Paginator(cobradores, 20)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'cobradores/cobrador_list.html', {
        'page_obj': page_obj,
        'query': query
    })


def cobrador_create(request):
    if request.method == 'POST':
        form = CobradorForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Cobrador creado exitosamente.')
            return redirect('cobradores:cobrador_list')  # ✅ Con namespace
    else:
        form = CobradorForm()
    return render(request, 'cobradores/cobrador_form.html', {
        'form': form,
        'title': 'Crear Cobrador'
    })


def cobrador_update(request, pk):
    cobrador = get_object_or_404(Cobrador, pk=pk)
    if request.method == 'POST':
        form = CobradorForm(request.POST, instance=cobrador)
        if form.is_valid():
            form.save()
            messages.success(request, 'Cobrador actualizado exitosamente.')
            return redirect('cobradores:cobrador_list')  # ✅ Con namespace
    else:
        form = CobradorForm(instance=cobrador)
    return render(request, 'cobradores/cobrador_form.html', {
        'form': form,
        'title': 'Editar Cobrador'
    })


@user_passes_test(es_admin_o_elegido)
def cobrador_delete(request, pk):
    cobrador = get_object_or_404(Cobrador, pk=pk)

    # ✅ Verificar si tiene movimientos
    tiene_cobros = Cobro.objects.filter(cobrador=cobrador).exists()
    tiene_devoluciones = Devolucion.objects.filter(cobrador=cobrador).exists()

    if tiene_cobros or tiene_devoluciones:
        messages.error(
            request,
            f'No se puede eliminar al cobrador "{cobrador.nombre}" porque tiene '
            f'{"pagos" if tiene_cobros else ""} '
            f'{"y " if tiene_cobros and tiene_devoluciones else ""}'
            f'{"devoluciones" if tiene_devoluciones else ""} registrados.'
        )
        return redirect('cobradores:cobrador_list')

    if request.method == 'POST':
        nombre = cobrador.nombre
        cobrador.delete()
        messages.success(request, f'Cobrador "{nombre}" eliminado exitosamente.')
        return redirect('cobradores:cobrador_list')

    return render(request, 'cobradores/cobrador_confirm_delete.html', {
        'cobrador': cobrador
    })


def cobrador_export_excel(request):
    cobradores = Cobrador.objects.all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cobradores"

    headers = ['Nombre', 'DNI', 'Teléfono', 'Correo', 'Dirección', 'Creado en']
    for col_num, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col_num)
        cell.value = header

    for cobrador in cobradores:
        row = [
            cobrador.nombre,
            cobrador.dni,
            cobrador.telefono or '',
            cobrador.correo or '',
            cobrador.direccion or '',
            cobrador.creado_en.strftime('%d/%m/%Y %H:%M')
        ]
        sheet.append(row)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=cobradores.xlsx'
    workbook.save(response)
    return response


def cobrador_export_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename=cobradores.csv'

    writer = csv.writer(response)
    writer.writerow(['Nombre', 'DNI', 'Teléfono', 'Correo', 'Dirección', 'Creado en'])

    for cobrador in Cobrador.objects.all():
        writer.writerow([
            cobrador.nombre,
            cobrador.dni,
            cobrador.telefono or '',
            cobrador.correo or '',
            cobrador.direccion or '',
            cobrador.creado_en.strftime('%d/%m/%Y %H:%M')
        ])

    return response